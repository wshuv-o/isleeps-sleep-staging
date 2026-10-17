"""Assemble the BSPC submission package from the files in this folder.

    python make_submission_package.py   ->  BSPC_submission/  and  BSPC_submission.zip

The manuscript PDF must already be compiled (mmnet_bspc.pdf) and the Overleaf zip current.
"""
import os
import shutil
import zipfile

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "BSPC_submission")
M = os.path.join(HERE, "manuscript")
FIGURES = [  # figure number in the manuscript -> source file
    (1, "figures/fig_mm_apnea_prev.pdf"), (2, "figures/fig_architecture.pdf"),
    (3, "figures/revision/fig_hypnogram_final.pdf"), (4, "figures/revision/fig_confusion_final.pdf"),
    (5, "figures/revision/fig_alignment_folds.pdf"), (6, "figures/revision/fig_rev_tsne.pdf"),
    (7, "figures/revision/fig_roc_pr_final.pdf"), (8, "figures/revision/fig_rev_ahi.pdf")]


def docx_from_lines(lines, path, bullets=False):
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    for i, line in enumerate(lines):
        if bullets and i == 0:
            doc.add_heading(line, level=1)
        elif bullets:
            doc.add_paragraph(line.lstrip("- ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(path)


def main():
    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    d = {k: os.path.join(OUT, k) for k in ("1_Manuscript", "2_LaTeX_source", "3_Highlights",
                                           "4_Graphical_abstract", "5_Cover_letter", "6_Figures",
                                           "7_Editable_figure_source")}
    for p in d.values():
        os.makedirs(p)

    shutil.copy(os.path.join(HERE, "mmnet_bspc.pdf"), os.path.join(d["1_Manuscript"], "MMNet_manuscript.pdf"))
    shutil.copy(os.path.join(HERE, "MMNet_BSPC_overleaf.zip"), os.path.join(d["2_LaTeX_source"], "MMNet_LaTeX_source.zip"))

    hl = [l.rstrip() for l in open(os.path.join(HERE, "highlights.txt"), encoding="utf-8") if l.strip()]
    shutil.copy(os.path.join(HERE, "highlights.txt"), os.path.join(d["3_Highlights"], "Highlights.txt"))
    docx_from_lines(hl, os.path.join(d["3_Highlights"], "Highlights.docx"), bullets=True)

    for ext in ("tif", "pdf", "png"):
        shutil.copy(os.path.join(HERE, "graphical_abstract.%s" % ext),
                    os.path.join(d["4_Graphical_abstract"], "Graphical_abstract.%s" % ext))

    cover = open(os.path.join(HERE, "cover_letter.txt"), encoding="utf-8").read()
    shutil.copy(os.path.join(HERE, "cover_letter.txt"), os.path.join(d["5_Cover_letter"], "Cover_letter.txt"))
    paras = [" ".join(p.split("\n")) if not p.lstrip().startswith("-") else p for p in cover.split("\n\n")]
    docx_from_lines([p.strip() for p in paras if p.strip()], os.path.join(d["5_Cover_letter"], "Cover_letter.docx"))

    for n, rel in FIGURES:
        shutil.copy(os.path.join(M, rel), os.path.join(d["6_Figures"], "Figure_%d.pdf" % n))
    shutil.copy(os.path.join(HERE, "fig_architecture.drawio"),
                os.path.join(d["7_Editable_figure_source"], "Figure_2_architecture.drawio"))

    zpath = OUT + ".zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(OUT):
            for f in files:
                p = os.path.join(root, f)
                z.write(p, os.path.relpath(p, os.path.dirname(OUT)))
    for root, _, files in os.walk(OUT):
        for f in sorted(files):
            print(os.path.relpath(os.path.join(root, f), OUT))
    print("->", zpath, "%.1f MB" % (os.path.getsize(zpath) / 1e6))


if __name__ == "__main__":
    main()
